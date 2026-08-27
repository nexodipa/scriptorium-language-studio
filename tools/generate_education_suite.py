from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
PORTFOLIO = ROOT / "portfolio"

INK = "172032"
BLUE = "476A78"
GOLD = "C9A857"
MUTED = "667386"
PALE = "F4F7F7"


DOCUMENTS = [
    {
        "slug": "scriptorium-plan-educativo-semanas-3-13",
        "title": "Plan educativo y comercial - semanas 3 a 13",
        "subtitle": "Hoja de ruta para convertir Scriptorium en un estudio lingüístico con una línea educativa real",
        "sections": [
            ("Objetivo", [
                "Construir dos líneas complementarias: servicios lingüísticos y educación. La meta de estas once semanas es pasar de una colección de muestras a una oferta pequeña, verificable y lista para recibir consultas y primeras ventas.",
                "La formalización tributaria y la facturación directa quedan fuera de esta etapa. Se revisarán antes de habilitar un checkout propio o emitir comprobantes desde la web.",
            ]),
            ("Semanas 3 a 5 - Evidencia y oferta", [
                "LIST:Publicar el alcance real de la formación lingüística y distinguir certificados de finalización de títulos oficiales.|Definir dos programas iniciales: Traducción y localización práctica e Introducción a la lectura de textos históricos.|Usar tres recursos terminados como herramientas de captación y validación.|Solicitar tres proyectos piloto o clases iniciales y documentar resultados con permiso.",
            ]),
            ("Semanas 6 a 9 - Autoridad y comunidad", [
                "LIST:Configurar Facebook, Instagram y YouTube con la misma identidad y descripción.|Publicar dos piezas por semana: una lección útil y una solución a un problema concreto.|Grabar tres videos iniciales: traducción vs. localización, cómo preparar un texto para traducir y cómo leer una fuente histórica sin exagerar certezas.|Registrar consultas, fuente de llegada y temas solicitados.",
            ]),
            ("Semanas 10 a 13 - Productos y conversión", [
                "LIST:Validar el brief de traducción, el checklist de localización y la guía de lectura histórica con usuarios reales.|Definir precio, licencia, soporte y política de actualización para cada producto.|Elegir Payhip, Ko-fi o Fiverr como canal inicial; Shopify se evalúa después de ventas recurrentes.|Conectar la web únicamente con productos realmente disponibles.",
            ]),
            ("Indicadores semanales", [
                "LIST:Consultas calificadas|Clases o proyectos reservados|Descargas de recursos|Respuestas y guardados en contenido educativo|Tiempo de respuesta|Conversión de consulta a venta|Comentarios que revelan nuevas necesidades",
            ]),
            ("Decisiones que requieren al fundador", [
                "LIST:Aprobar precios y duración de clases.|Confirmar niveles lingüísticos que pueden publicarse.|Autorizar qué certificados pueden mostrarse con imagen o enlace.|Elegir plataforma de cobro cuando los productos hayan sido probados.|Participar en la creación final de cuentas sociales y cualquier acción de publicación.",
            ]),
        ],
    },
    {
        "slug": "scriptorium-brief-traduccion-profesional",
        "title": "Brief profesional de traducción",
        "subtitle": "Plantilla editable para cotizar, traducir y revisar con un alcance claro",
        "sections": [
            ("1. Datos del proyecto", [
                "FIELDS:Nombre del proyecto|Persona o institución|Correo de contacto|Fecha de solicitud|Fecha de entrega deseada",
            ]),
            ("2. Idiomas y contenido", [
                "FIELDS:Idioma de origen|Idioma de destino|Cantidad aproximada de palabras|Tipo de documento|Formato de entrega",
                "CHECKS:Texto editable disponible|Incluye imágenes con texto|Incluye tablas o referencias|Requiere mantener diseño|Contiene información sensible",
            ]),
            ("3. Propósito y lector", [
                "FIELDS:¿Dónde se publicará o utilizará?|¿Quién leerá el texto?|Tono esperado|País o variante lingüística|Términos que deben conservarse",
            ]),
            ("4. Referencias", [
                "FIELDS:Glosario disponible|Guía de estilo|Enlaces o documentos de referencia|Traducciones anteriores aprobadas|Persona que validará el contenido",
            ]),
            ("5. Alcance acordado", [
                "CHECKS:Traducción|Localización cultural|Edición del texto origen|Revisión bilingüe|Maquetación básica|Control de calidad final",
                "FIELDS:Número de revisiones incluidas|Canal de entrega|Observaciones y exclusiones",
            ]),
            ("Nota de uso", [
                "Esta plantilla organiza la cotización, pero no sustituye un acuerdo de servicios. No envíes documentos médicos, legales o personales sensibles hasta acordar un canal y un tratamiento adecuado.",
            ]),
        ],
    },
    {
        "slug": "scriptorium-checklist-localizacion-web",
        "title": "Checklist de localización web",
        "subtitle": "Control práctico para adaptar una web a otro idioma sin romper su experiencia",
        "sections": [
            ("Antes de traducir", [
                "CHECKS:Definir idioma, país y audiencia|Confirmar páginas incluidas|Extraer textos de botones, formularios y mensajes|Reunir glosario y tono de marca|Guardar una copia de referencia",
            ]),
            ("Contenido y significado", [
                "CHECKS:Adaptar expresiones, no traducir literalmente|Revisar nombres de servicios y llamadas a la acción|Comprobar fechas, monedas, medidas y direcciones|Mantener terminología consistente|Marcar textos que requieren aprobación experta",
            ]),
            ("Interfaz", [
                "CHECKS:Verificar que botones y menús no se corten|Revisar formularios, errores y confirmaciones|Comprobar texto alternativo de imágenes|Revisar títulos, metadescripciones y datos estructurados|Validar enlaces y descargas",
            ]),
            ("Prueba final", [
                "CHECKS:Revisar en computadora y móvil|Probar navegación completa|Buscar fragmentos que quedaron en el idioma original|Validar tipografías y caracteres|Obtener una segunda lectura lingüística cuando sea posible|Registrar versión y fecha de entrega",
            ]),
            ("Resultado", [
                "FIELDS:Páginas revisadas|Incidencias encontradas|Incidencias corregidas|Pendientes del cliente|Persona que aprobó|Fecha y versión",
            ]),
        ],
    },
    {
        "slug": "scriptorium-guia-lectura-textos-historicos",
        "title": "Guía inicial para leer textos históricos",
        "subtitle": "Un método prudente para separar escritura, transliteración, traducción y contexto",
        "sections": [
            ("Antes de interpretar", [
                "Una fuente histórica no empieza con una traducción definitiva. Primero se identifica su procedencia, soporte, fecha aproximada, sistema de escritura y estado de conservación. Una imagen aislada rara vez basta para una conclusión firme.",
            ]),
            ("Las cinco capas", [
                "LIST:Fuente: de dónde procede la imagen o edición.|Original: qué signos o letras aparecen realmente.|Transliteración: cómo se representan esos signos en otro alfabeto.|Traducción: una propuesta de significado, con alternativas cuando existan.|Contexto: género, época, función, audiencia y límites de la lectura.",
            ]),
            ("Preguntas de control", [
                "CHECKS:¿La fuente es primaria, una fotografía o una edición moderna?|¿Hay signos dañados o reconstruidos?|¿La transliteración sigue una convención identificada?|¿La traducción distingue certeza de hipótesis?|¿Se citan edición, autor o repositorio?|¿La explicación evita convertir una lectura posible en verdad absoluta?",
            ]),
            ("Ficha de análisis", [
                "FIELDS:Título o identificación de la fuente|Repositorio o referencia|Fecha y lugar aproximados|Sistema de escritura|Dirección de lectura|Transliteración utilizada|Traducción propuesta|Alternativas|Contexto|Dudas pendientes",
            ]),
            ("Alcance", [
                "Esta guía es introductoria. El análisis académico, epigráfico, paleográfico, jurídico o religioso requiere bibliografía especializada y, según el caso, revisión de un profesional competente.",
            ]),
        ],
    },
    {
        "slug": "scriptorium-programa-traduccion-localizacion",
        "title": "Programa inicial: traducción y localización práctica",
        "subtitle": "Cuatro módulos para comprender el proceso y producir una primera entrega revisada",
        "sections": [
            ("Propósito", [
                "Programa introductorio para estudiantes, creadores y pequeños negocios. Al finalizar, la persona podrá preparar un brief, distinguir traducción de localización, crear un glosario básico y revisar una pieza web breve.",
            ]),
            ("Módulo 1 - Propósito, lector y contexto", [
                "LIST:Qué cambia cuando cambia el lector.|Traducción, localización, transcreación y edición.|Ejercicio: detectar decisiones invisibles en dos versiones de una página.",
            ]),
            ("Módulo 2 - Brief y glosario", [
                "LIST:Cómo pedir la información correcta.|Términos preferidos, prohibidos y variables.|Ejercicio: completar el brief de un proyecto real o simulado.",
            ]),
            ("Módulo 3 - Adaptación web", [
                "LIST:Botones, menús, formularios, errores y metadatos.|Longitud, tono y consistencia.|Ejercicio: localizar una página breve con el checklist.",
            ]),
            ("Módulo 4 - Revisión y entrega", [
                "LIST:Lectura bilingüe y lectura natural.|Control de números, nombres, enlaces y formato.|Ejercicio final: entregar versión, glosario y notas de decisión.",
            ]),
            ("Modalidad sugerida", [
                "LIST:Cuatro sesiones de 60 minutos.|Actividad breve entre sesiones.|Materiales incluidos: brief, checklist y guía de revisión.|Evaluación formativa; no otorga una certificación oficial.",
            ]),
        ],
    },
    {
        "slug": "scriptorium-plan-contenidos-30-dias",
        "title": "Plan de contenidos - primeros 30 días",
        "subtitle": "Facebook, Instagram y YouTube con una voz educativa, útil y verificable",
        "sections": [
            ("Cadencia", [
                "Publicar dos piezas por semana y un video largo cada diez días. Cada publicación debe resolver una pregunta real, citar fuentes cuando incluya datos y cerrar con una sola acción clara.",
            ]),
            ("Calendario", [
                "LIST:Semana 1: traducción vs. localización + carrusel del brief.|Semana 2: cinco errores al traducir una web + video sobre cómo preparar archivos.|Semana 3: original, transliteración y traducción + muestra de la guía histórica.|Semana 4: caso práctico de botón, formulario y tono + invitación a clase piloto.",
            ]),
            ("Video 1 - Traducir no es reemplazar palabras", [
                "Apertura: dos frases pueden significar lo mismo y funcionar distinto. Desarrollo: objetivo, lector, tono, formato y cultura. Demostración: una llamada a la acción literal frente a una versión localizada. Cierre: descarga el brief antes de cotizar tu próximo proyecto.",
            ]),
            ("Video 2 - Cómo preparar una web para traducir", [
                "Apertura: el texto visible es solo una parte de la web. Desarrollo: menús, botones, formularios, errores, metadatos e imágenes. Demostración: recorrer el checklist. Cierre: solicita una revisión de localización.",
            ]),
            ("Video 3 - Cómo acercarse a un texto histórico", [
                "Apertura: una inscripción no trae una traducción única pegada debajo. Desarrollo: fuente, original, transliteración, traducción y contexto. Demostración: completar una ficha sin fingir certeza. Cierre: revisa la lección introductoria.",
            ]),
            ("Descripción del canal", [
                "Scriptorium Language Studio publica lecciones breves sobre traducción, localización, filología aplicada, escritura y lectura responsable de textos históricos. Contenido educativo en español, con ejemplos, fuentes y límites visibles.",
            ]),
        ],
    },
]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def configure_doc(doc: Document, title: str, subtitle: str):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in [(1, 16), (2, 13), (3, 11.5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "SCRIPTORIUM LANGUAGE STUDIO  |  EDUCACIÓN Y FILOLOGÍA APLICADA"
    header.runs[0].font.name = "Calibri"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Scriptorium - material educativo | josuepug@gmail.com")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("SCRIPTORIUM | RECURSO EDUCATIVO")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(GOLD)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    r.font.name = "Georgia"
    r.font.size = Pt(25)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)
    s = doc.add_paragraph(subtitle)
    s.paragraph_format.space_after = Pt(14)
    s.runs[0].font.name = "Calibri"
    s.runs[0].font.size = Pt(11.5)
    s.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def add_fields(doc: Document, fields: list[str]):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.3)
    for i, label in enumerate(fields):
        row = table.rows[0] if i == 0 else table.add_row()
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(4.3)
        row.cells[0].text = label
        row.cells[1].text = " "
        set_cell_shading(row.cells[0], PALE)
        row.cells[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()


def add_checks(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("[  ] " + item)


def add_sections_docx(doc: Document, sections):
    for heading, blocks in sections:
        doc.add_heading(heading, level=1)
        for block in blocks:
            if block.startswith("LIST:"):
                for item in block[5:].split("|"):
                    doc.add_paragraph(item, style="List Bullet")
            elif block.startswith("CHECKS:"):
                add_checks(doc, block[7:].split("|"))
            elif block.startswith("FIELDS:"):
                add_fields(doc, block[7:].split("|"))
            else:
                doc.add_paragraph(block)


def build_docx(spec):
    doc = Document()
    configure_doc(doc, spec["title"], spec["subtitle"])
    add_sections_docx(doc, spec["sections"])
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = "Scriptorium Language Studio"
    out = PORTFOLIO / f'{spec["slug"]}.docx'
    doc.save(out)
    return out


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "kicker": ParagraphStyle("Kicker", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#" + GOLD), spaceAfter=8),
        "title": ParagraphStyle("Title", parent=styles["Title"], fontName="Times-Bold", fontSize=21, leading=24, textColor=colors.HexColor("#" + INK), alignment=TA_LEFT, spaceAfter=5),
        "subtitle": ParagraphStyle("Subtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=10.5, leading=14, textColor=colors.HexColor("#" + MUTED), spaceAfter=16),
        "h1": ParagraphStyle("H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=12.8, leading=15, textColor=colors.HexColor("#" + BLUE), spaceBefore=8, spaceAfter=5),
        "body": ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12.5, textColor=colors.HexColor("#" + INK), spaceAfter=5),
        "bullet": ParagraphStyle("Bullet", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12.5, leftIndent=14, firstLineIndent=-8, textColor=colors.HexColor("#" + INK), spaceAfter=3.5),
        "small": ParagraphStyle("Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#" + MUTED)),
    }


def pdf_header_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D7E0E4"))
    canvas.line(0.75 * inch, letter[1] - 0.55 * inch, letter[0] - 0.75 * inch, letter[1] - 0.55 * inch)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#" + MUTED))
    canvas.drawString(0.75 * inch, letter[1] - 0.43 * inch, "SCRIPTORIUM LANGUAGE STUDIO | EDUCACIÓN Y FILOLOGÍA APLICADA")
    canvas.drawString(0.75 * inch, 0.4 * inch, "Material educativo - alcance introductorio")
    canvas.drawRightString(letter[0] - 0.75 * inch, 0.4 * inch, f"Página {doc.page}")
    canvas.restoreState()


def build_pdf(spec):
    styles = pdf_styles()
    out = PORTFOLIO / f'{spec["slug"]}.pdf'
    doc = SimpleDocTemplate(str(out), pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.75 * inch, bottomMargin=0.65 * inch, title=spec["title"], author="Scriptorium Language Studio")
    story = [
        Paragraph("SCRIPTORIUM | RECURSO EDUCATIVO", styles["kicker"]),
        Paragraph(spec["title"], styles["title"]),
        Paragraph(spec["subtitle"], styles["subtitle"]),
    ]
    for heading, blocks in spec["sections"]:
        if spec["slug"] == "scriptorium-plan-educativo-semanas-3-13" and heading.startswith("Semanas 10"):
            story.append(PageBreak())
        if spec["slug"] == "scriptorium-programa-traduccion-localizacion" and heading.startswith("Módulo 3"):
            story.append(PageBreak())
        if spec["slug"] == "scriptorium-checklist-localizacion-web" and heading == "Resultado":
            story.append(PageBreak())
        if spec["slug"] == "scriptorium-guia-lectura-textos-historicos" and heading == "Ficha de análisis":
            story.append(PageBreak())
        story.append(Paragraph(heading, styles["h1"]))
        for block in blocks:
            if block.startswith("LIST:"):
                for item in block[5:].split("|"):
                    story.append(Paragraph("• " + item, styles["bullet"]))
            elif block.startswith("CHECKS:"):
                for item in block[7:].split("|"):
                    story.append(Paragraph("[ ] " + item, styles["bullet"]))
            elif block.startswith("FIELDS:"):
                rows = [[Paragraph(label, styles["small"]), ""] for label in block[7:].split("|")]
                table = Table(rows, colWidths=[1.8 * inch, 4.45 * inch], rowHeights=[0.29 * inch] * len(rows))
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#" + PALE)),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D7E0E4")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ]))
                story.extend([table, Spacer(1, 8)])
            else:
                story.append(Paragraph(block, styles["body"]))
    doc.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=["docx", "pdf", "all"], default="all")
    args = parser.parse_args()
    PORTFOLIO.mkdir(exist_ok=True)
    for spec in DOCUMENTS:
        if args.format in ("docx", "all"):
            print(build_docx(spec).relative_to(ROOT))
        if args.format in ("pdf", "all"):
            print(build_pdf(spec).relative_to(ROOT))


if __name__ == "__main__":
    main()
