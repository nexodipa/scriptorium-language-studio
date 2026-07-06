from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
PORTFOLIO = ROOT / "portfolio"

INK = "#172032"
BLUE_DEEP = "#243f50"
BLUE = "#476a78"
MUTED = "#667386"
GOLD = "#c9a857"
IVORY = "#fbf7ee"
LINE = "#d7e0e4"


def font(kind: str, size: int):
    choices = {
        "serif": [r"C:\Windows\Fonts\georgia.ttf", r"C:\Windows\Fonts\times.ttf"],
        "serif-bold": [r"C:\Windows\Fonts\georgiab.ttf", r"C:\Windows\Fonts\timesbd.ttf"],
        "sans": [r"C:\Windows\Fonts\segoeui.ttf", r"C:\Windows\Fonts\arial.ttf"],
        "sans-bold": [r"C:\Windows\Fonts\segoeuib.ttf", r"C:\Windows\Fonts\arialbd.ttf"],
    }
    for path in choices[kind]:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, max_width: int, face, max_lines: int | None = None) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=face)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    if max_lines and len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip(".") + "..."
    return lines


def draw_wrapped(draw: ImageDraw.ImageDraw, text: str, x: int, y: int, max_width: int, face, fill: str, gap: int) -> int:
    for line in wrap(draw, text, max_width, face):
        draw.text((x, y), line, font=face, fill=fill)
        bbox = draw.textbbox((x, y), line, font=face)
        y += (bbox[3] - bbox[1]) + gap
    return y


def cover_crop(path: Path, size: tuple[int, int], anchor: str = "center") -> Image.Image:
    img = Image.open(path).convert("RGB")
    tw, th = size
    scale = max(tw / img.width, th / img.height)
    img = img.resize((int(img.width * scale), int(img.height * scale)), Image.Resampling.LANCZOS)
    left = (img.width - tw) // 2
    if anchor == "left":
        left = 0
    elif anchor == "right":
        left = img.width - tw
    top = (img.height - th) // 2
    return img.crop((left, top, left + tw, top + th))


def make_social_post(filename: str, background: str, eyebrow: str, title: str, body: str, footer: str, anchor: str = "center") -> Path:
    size = (1080, 1080)
    base = cover_crop(ASSETS / background, size, anchor).convert("RGBA")
    overlay = Image.new("RGBA", size, (251, 247, 238, 0))
    od = ImageDraw.Draw(overlay)
    od.rectangle((0, 0, 1080, 1080), fill=(251, 247, 238, 208))
    od.rectangle((0, 0, 1080, 1080), outline=(201, 168, 87, 150), width=6)
    base = Image.alpha_composite(base, overlay)
    draw = ImageDraw.Draw(base)
    logo = Image.open(ASSETS / "logo-icon.png").convert("RGBA").resize((84, 84), Image.Resampling.LANCZOS)
    base.alpha_composite(logo, (80, 76))
    draw.text((80, 190), "SCRIPTORIUM", font=font("serif-bold", 50), fill=INK)
    draw.text((80, 260), eyebrow.upper(), font=font("sans-bold", 22), fill=BLUE_DEEP)
    y = draw_wrapped(draw, title, 80, 350, 820, font("serif-bold", 66), INK, 10)
    y += 24
    draw_wrapped(draw, body, 80, y, 760, font("sans", 32), BLUE_DEEP, 12)
    draw.line((80, 938, 1000, 938), fill=GOLD, width=4)
    draw.text((80, 968), footer, font=font("sans-bold", 22), fill=BLUE_DEEP)
    out = ASSETS / filename
    base.convert("RGB").save(out, quality=94)
    return out


def make_course_card(filename: str, background: str, title: str, meta: str, bullets: list[str], anchor: str = "center") -> Path:
    card = cover_crop(ASSETS / background, (1400, 900), anchor).convert("RGBA")
    overlay = Image.new("RGBA", card.size, (255, 255, 255, 0))
    od = ImageDraw.Draw(overlay)
    od.rectangle((0, 0, 850, 900), fill=(255, 253, 247, 236))
    card = Image.alpha_composite(card, overlay)
    draw = ImageDraw.Draw(card)
    draw.rectangle((70, 70, 1330, 830), outline=(201, 168, 87, 160), width=3)
    draw.text((118, 120), "SCRIPTORIUM", font=font("serif-bold", 36), fill=INK)
    draw.text((118, 178), meta.upper(), font=font("sans-bold", 18), fill=GOLD)
    draw_wrapped(draw, title, 118, 270, 620, font("serif-bold", 56), INK, 8)
    y = 520
    for bullet in bullets:
        draw.text((130, y), "•", font=font("serif-bold", 28), fill=GOLD)
        y = draw_wrapped(draw, bullet, 170, y, 560, font("sans", 27), BLUE_DEEP, 8) + 10
    out = ASSETS / filename
    card.convert("RGB").save(out, quality=94)
    return out


def doc_setup(doc: Document, title: str):
    section = doc.sections[0]
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, attr, Inches(1))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor.from_string("172032")
    style.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.name = "Georgia"
    r.font.size = Pt(24)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("172032")
    subtitle = doc.add_paragraph("Scriptorium - cursos, clases y materiales pedagógicos")
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("667386")


def build_docx(filename: str, title: str, sections: list[tuple[str, list[str]]], image: str | None = None) -> Path:
    doc = Document()
    doc_setup(doc, title)
    if image:
        doc.add_picture(str(ASSETS / image), width=Inches(6.5))
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for paragraph in paragraphs:
            if paragraph.startswith("LIST:"):
                for item in paragraph[5:].split("|"):
                    doc.add_paragraph(item, style="List Bullet")
            else:
                doc.add_paragraph(paragraph)
    out = PORTFOLIO / f"{filename}.docx"
    doc.save(out)
    return out


def build_pdf(filename: str, title: str, sections: list[tuple[str, list[str]]]) -> Path:
    out = PORTFOLIO / f"{filename}.pdf"
    c = canvas.Canvas(str(out), pagesize=letter)
    w, h = letter
    x = inch
    y = h - inch
    c.setFillColor(colors.HexColor(INK))
    c.setFont("Times-Bold", 23)
    c.drawString(x, y, "Scriptorium")
    y -= 34
    c.setFillColor(colors.HexColor(GOLD))
    c.setFont("Helvetica-Bold", 8)
    c.drawString(x, y, "CURSOS, CLASES Y MATERIALES PEDAGÓGICOS")
    y -= 42
    c.setFillColor(colors.HexColor(INK))
    c.setFont("Times-Bold", 20)
    c.drawString(x, y, title[:80])
    y -= 38
    for heading, paragraphs in sections:
        if y < 1.2 * inch:
            c.showPage()
            y = h - inch
        c.setFillColor(colors.HexColor(BLUE_DEEP))
        c.setFont("Helvetica-Bold", 13)
        c.drawString(x, y, heading)
        y -= 20
        c.setFillColor(colors.HexColor(MUTED))
        c.setFont("Helvetica", 9.6)
        for paragraph in paragraphs:
            items = paragraph[5:].split("|") if paragraph.startswith("LIST:") else [paragraph]
            for item in items:
                prefix = "- " if paragraph.startswith("LIST:") else ""
                words = (prefix + item).split()
                line = ""
                for word in words:
                    trial = f"{line} {word}".strip()
                    if c.stringWidth(trial, "Helvetica", 9.6) <= w - 2 * inch:
                        line = trial
                    else:
                        c.drawString(x, y, line)
                        y -= 14
                        line = word
                if line:
                    c.drawString(x, y, line)
                    y -= 16
            y -= 4
        y -= 10
    c.setFillColor(colors.HexColor(MUTED))
    c.setFont("Helvetica", 8)
    c.drawRightString(w - inch, 0.45 * inch, "Documento demo - Scriptorium")
    c.save()
    PdfReader(str(out))
    return out


def main():
    PORTFOLIO.mkdir(exist_ok=True)
    created = []
    created.extend([
        make_social_post(
            "post-translation-localization.png",
            "translation-workspace.png",
            "Traducción profesional",
            "Traducir no es copiar palabras",
            "El texto se adapta a objetivo, lector, formato y tono. Por eso cada proyecto se cotiza por alcance.",
            "Guarda este post si vas a traducir una web, tesis o documento.",
            "right",
        ),
        make_social_post(
            "post-ancient-languages.png",
            "ancient-language-classroom.png",
            "Lenguas antiguas",
            "Leer un texto antiguo exige contexto",
            "Original, transliteración, traducción y nota cultural deben verse separados para no confundir lectura con interpretación.",
            "Latín · Griego · Hebreo · Árabe clásico · Sánscrito",
            "center",
        ),
        make_social_post(
            "post-course-launch.png",
            "pedagogical-shop.png",
            "Cursos y materiales",
            "Aprende idiomas con documentos reales",
            "Clases, mini cursos y guías descargables para traducción, localización y primeras lecturas de lenguas históricas.",
            "Solicita clase o material por WhatsApp/Fiverr.",
            "center",
        ),
        make_social_post(
            "post-translation-brief.png",
            "translation-workspace.png",
            "Antes de cotizar",
            "Un buen brief ahorra tiempo y errores",
            "Envía idioma origen, idioma destino, extensión, formato, fecha, uso final y nivel de confidencialidad.",
            "Plantilla disponible en la web.",
            "left",
        ),
        make_course_card(
            "course-card-translation-localization.png",
            "translation-workspace.png",
            "Curso práctico de traducción y localización",
            "Curso guiado",
            ["4 módulos introductorios", "Ejercicios con textos web y académicos", "Entrega de guía PDF y checklist"],
            "right",
        ),
        make_course_card(
            "course-card-ancient-languages.png",
            "ancient-language-classroom.png",
            "Introducción a lenguas antiguas",
            "Clase / mini curso",
            ["Latín, griego y hebreo como ruta inicial", "Transliteración y notas culturales", "Material pedagógico descargable"],
            "center",
        ),
        make_course_card(
            "course-card-pedagogical-shop.png",
            "pedagogical-shop.png",
            "Tienda de documentos pedagógicos",
            "Recursos descargables",
            ["Fichas, guías y glosarios", "PDF y DOCX editables", "Material demo antes de compra o encargo"],
            "center",
        ),
    ])

    docs = [
        (
            "scriptorium-course-catalog",
            "Catálogo de cursos y clases",
            "course-card-translation-localization.png",
            [
                ("Ruta de traducción y localización", [
                    "Curso pensado para estudiantes, creadores y pequeños negocios que necesitan entender cómo se adapta un texto a otro idioma sin perder tono ni objetivo.",
                    "LIST:Módulo 1: traducción vs localización|Módulo 2: brief, público y tono|Módulo 3: glosario y consistencia|Módulo 4: revisión y entrega final",
                ]),
                ("Ruta de lenguas antiguas", [
                    "Clases introductorias para leer muestras breves con método: escritura, transliteración, traducción tentativa y nota cultural.",
                    "LIST:Latín introductorio|Griego antiguo para citas y formas básicas|Hebreo bíblico con cuidado de dirección y transliteración|Rúnico como muestra histórica y material",
                ]),
                ("Modalidad", [
                    "Clases por cotización, sesiones individuales o pequeños grupos, con entrega de material PDF/DOCX cuando corresponda. No sustituye formación universitaria ni certificación oficial.",
                ]),
            ],
        ),
        (
            "scriptorium-pedagogical-store-catalog",
            "Catálogo de tienda pedagógica",
            "course-card-pedagogical-shop.png",
            [
                ("Productos iniciales", [
                    "LIST:Brief de traducción editable|Checklist de localización web|Fichas de alfabeto y transliteración|Glosario académico básico|Mini guía de lectura de textos antiguos",
                ]),
                ("Cómo se vendería", [
                    "Cada producto puede ofrecerse como descarga demo, pack completo por encargo o material incluido en clases. En la web actual se solicita por WhatsApp/Fiverr hasta activar pagos propios.",
                ]),
                ("Diferenciador", [
                    "La tienda no debe parecer un banco de PDFs sueltos: cada material debe explicar para quién es, qué problema resuelve, qué incluye y cómo se usa.",
                ]),
            ],
        ),
        (
            "scriptorium-ancient-languages-sample-lesson",
            "Lección demo: cómo leer una muestra antigua",
            "course-card-ancient-languages.png",
            [
                ("Objetivo de la lección", [
                    "Aprender a mirar un fragmento antiguo sin saltar directo a una traducción absoluta. Primero se separa soporte, escritura, transliteración, vocabulario, contexto y límite de interpretación.",
                ]),
                ("Estructura de clase", [
                    "LIST:Observación visual del fragmento|Identificación de escritura o alfabeto|Transliteración normalizada cuando aplique|Traducción de muestra|Nota cultural y advertencia de alcance",
                ]),
                ("Resultado esperado", [
                    "El estudiante comprende por qué un texto antiguo exige prudencia y por qué un buen documento pedagógico distingue lectura, hipótesis y explicación.",
                ]),
            ],
        ),
    ]
    for slug, title, image, sections in docs:
        created.append(build_docx(slug, title, sections, image))
        created.append(build_pdf(slug, title, sections))
    for path in created:
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
