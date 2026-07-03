from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)


ROOT = Path(__file__).resolve().parents[1]
PORTFOLIO = ROOT / "portfolio"
ASSETS = ROOT / "assets"
PORTFOLIO.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

BLUE = "#476A78"
BLUE_DARK = "#1F3A44"
GOLD = "#B8932F"
IVORY = "#F7F4EB"
INK = "#24313A"
MUTED = "#5C6B72"


def docx_style(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, BLUE_DARK),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.text = "Scriptorium - professional sample, non-certified language and documentation work"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(92, 107, 114)


def add_title(doc, title, subtitle, metadata):
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title_p.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE_DARK.replace("#", ""))

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    r = sub.add_run(subtitle)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(92, 107, 114)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    return table


def save_docx(filename, title, subtitle, metadata, sections):
    doc = Document()
    docx_style(doc)
    add_title(doc, title, subtitle, metadata)
    for section in sections:
        kind = section[0]
        if kind == "h1":
            doc.add_heading(section[1], level=1)
        elif kind == "h2":
            doc.add_heading(section[1], level=2)
        elif kind == "p":
            doc.add_paragraph(section[1])
        elif kind == "bullets":
            add_bullets(doc, section[1])
        elif kind == "table":
            add_table(doc, section[1], section[2])
        elif kind == "break":
            doc.add_page_break()
    path = PORTFOLIO / filename
    doc.save(path)
    return path


def pdf_styles():
    st = getSampleStyleSheet()
    st.add(ParagraphStyle("TitleS", parent=st["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor(BLUE_DARK), spaceAfter=8))
    st.add(ParagraphStyle("Sub", parent=st["BodyText"], fontName="Helvetica", fontSize=10, leading=14, textColor=colors.HexColor(MUTED), spaceAfter=14))
    st.add(ParagraphStyle("H1S", parent=st["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.HexColor(BLUE), spaceBefore=14, spaceAfter=8))
    st.add(ParagraphStyle("H2S", parent=st["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor(BLUE_DARK), spaceBefore=10, spaceAfter=5))
    st.add(ParagraphStyle("BodyS", parent=st["BodyText"], fontName="Helvetica", fontSize=9.4, leading=13, textColor=colors.HexColor(INK), spaceAfter=7))
    st.add(ParagraphStyle("Small", parent=st["BodyText"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor(MUTED)))
    return st


def on_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D7E0E4"))
    canvas.line(0.72 * inch, 0.58 * inch, 7.78 * inch, 0.58 * inch)
    canvas.setFont("Helvetica", 7.4)
    canvas.setFillColor(colors.HexColor(MUTED))
    canvas.drawString(0.72 * inch, 0.42 * inch, "Scriptorium - professional sample, non-certified language and documentation work")
    canvas.drawRightString(7.78 * inch, 0.42 * inch, f"Page {doc.page}")
    canvas.restoreState()


def save_pdf(filename, title, subtitle, metadata, sections):
    st = pdf_styles()
    story = [Paragraph(title, st["TitleS"]), Paragraph(subtitle, st["Sub"])]
    data = [[Paragraph(f"<b>{a}</b>", st["Small"]), Paragraph(b, st["Small"])] for a, b in metadata]
    table = Table(data, colWidths=[1.6 * inch, 4.8 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(IVORY)),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7E0E4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [table, Spacer(1, 0.14 * inch)]
    for section in sections:
        kind = section[0]
        if kind == "h1":
            story.append(Paragraph(section[1], st["H1S"]))
        elif kind == "h2":
            story.append(Paragraph(section[1], st["H2S"]))
        elif kind == "p":
            story.append(Paragraph(section[1], st["BodyS"]))
        elif kind == "bullets":
            for item in section[1]:
                story.append(Paragraph(f"&bull; {item}", st["BodyS"]))
        elif kind == "table":
            rows = [[Paragraph(str(cell), st["Small"]) for cell in section[1]]]
            rows += [[Paragraph(str(cell), st["Small"]) for cell in row] for row in section[2]]
            col_width = 6.4 * inch / len(section[1])
            t = Table(rows, colWidths=[col_width] * len(section[1]), repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(BLUE_DARK)),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7E0E4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story += [Spacer(1, 0.08 * inch), t, Spacer(1, 0.1 * inch)]
        elif kind == "break":
            story.append(PageBreak())
    path = PORTFOLIO / filename
    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.72 * inch, bottomMargin=0.72 * inch)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return path


def create_cover(filename, title, subtitle, code):
    img = Image.new("RGB", (1200, 1500), "#DCE8EB")
    draw = ImageDraw.Draw(img)
    for y in range(1500):
        ratio = y / 1500
        r = int(247 * (1 - ratio) + 220 * ratio)
        g = int(244 * (1 - ratio) + 232 * ratio)
        b = int(235 * (1 - ratio) + 235 * ratio)
        draw.line((0, y, 1200, y), fill=(r, g, b))
    draw.rectangle((86, 90, 1114, 1410), outline="#B8932F", width=4)
    draw.rectangle((118, 122, 1082, 1378), outline="#FFFFFF", width=2)
    try:
        title_font = ImageFont.truetype("C:/Windows/Fonts/georgiab.ttf", 72)
        sub_font = ImageFont.truetype("C:/Windows/Fonts/calibri.ttf", 32)
        code_font = ImageFont.truetype("C:/Windows/Fonts/georgiab.ttf", 130)
    except OSError:
        title_font = sub_font = code_font = ImageFont.load_default()
    draw.text((170, 250), "SCRIPTORIUM", fill=BLUE_DARK, font=title_font)
    draw.line((170, 350, 1030, 350), fill=GOLD, width=3)
    draw.text((170, 405), title, fill=INK, font=title_font)
    draw.multiline_text((170, 590), subtitle, fill=MUTED, font=sub_font, spacing=10)
    draw.text((170, 1040), code, fill=GOLD, font=code_font)
    draw.text((170, 1230), "Professional portfolio sample", fill=BLUE_DARK, font=sub_font)
    path = ASSETS / filename
    img.save(path)
    return path


def build_all():
    languages = "Spanish, English, German, French, Portuguese, Italian, Russian, Czech, Chinese, Japanese, Hebrew and Arabic; historical samples in Latin, Ancient Greek, Classical Arabic, Sanskrit and Old Norse."

    samples = [
        {
            "slug": "scriptorium-language-access-brief",
            "title": "Language Access and Digital Trust",
            "subtitle": "Original bilingual research brief for service pages, onboarding and client communication.",
            "cover_title": "Language Access",
            "cover_subtitle": "Research brief and multilingual sample",
            "code": "LA",
            "metadata": [
                ("Document type", "Original research-style portfolio sample"),
                ("Use", "Academic, public information and service communication"),
                ("Commercial status", "Quote-only, non-certified language work"),
                ("Languages", languages),
            ],
            "sections": [
                ("h1", "Executive summary"),
                ("p", "Clear language access improves trust when a digital service asks users to share personal, academic or professional information. This original Scriptorium sample demonstrates how a short research brief can become a practical multilingual asset without pretending to be an external client case."),
                ("p", "The key finding is simple: users understand a service faster when the page separates the promise, the scope, the limits and the next action. Translation alone is not enough when the reader also needs orientation, safety and confidence."),
                ("h1", "Method used for this sample"),
                ("bullets", [
                    "A Spanish base text was drafted as a service-information brief.",
                    "The content was adapted into selected modern languages with attention to tone and usage.",
                    "Sensitive claims were removed: no certification, no legal advice and no invented client outcome.",
                    "Terminology was kept stable: access, scope, confidentiality, revision and delivery.",
                ]),
                ("h1", "Multilingual excerpt"),
                ("table", ["Language", "Localized sample"], [
                    ["Spanish", "La claridad lingüística reduce fricción cuando el usuario debe compartir información sensible."],
                    ["English", "Clear language reduces friction when users need to share sensitive information."],
                    ["German", "Klare Sprache verringert Reibung, wenn Nutzer sensible Informationen teilen müssen."],
                    ["French", "Un langage clair réduit la friction lorsque l'utilisateur doit partager des informations sensibles."],
                    ["Portuguese", "A linguagem clara reduz atritos quando o usuário precisa compartilhar informações sensíveis."],
                    ["Italian", "Un linguaggio chiaro riduce l'attrito quando l'utente deve condividere informazioni sensibili."],
                    ["Russian", "Ясный язык снижает барьер, когда пользователь должен передать чувствительную информацию."],
                    ["Czech", "Jasný jazyk snižuje tření, když má uživatel sdílet citlivé informace."],
                    ["Chinese", "清晰的语言能降低用户分享敏感信息时的阻力。"],
                    ["Japanese", "明確な表現は、利用者が機密情報を共有する際の不安を減らします。"],
                    ["Hebrew", "שפה ברורה מפחיתה חיכוך כאשר משתמשים צריכים לשתף מידע רגיש."],
                    ["Arabic", "تقلل اللغة الواضحة من التردد عندما يحتاج المستخدم إلى مشاركة معلومات حساسة."],
                ]),
                ("h1", "Client-ready use"),
                ("p", "This sample can be used to show how Scriptorium prepares service pages, onboarding notes, academic summaries and public-facing information with multilingual clarity."),
            ],
        },
        {
            "slug": "scriptorium-academic-abstract-pack",
            "title": "Academic Abstract and Terminology Pack",
            "subtitle": "Original academic sample with abstract adaptation, terminology control and editorial notes.",
            "cover_title": "Academic Abstract",
            "cover_subtitle": "Research summary, glossary and editorial notes",
            "code": "AA",
            "metadata": [
                ("Document type", "Original academic documentation sample"),
                ("Use", "Thesis, article abstract, literature review and research proposal support"),
                ("Commercial status", "Quote-only, non-certified language and editorial support"),
                ("Languages", languages),
            ],
            "sections": [
                ("h1", "Base abstract"),
                ("p", "This sample examines how bilingual service documentation can increase comprehension in small organizations that work with international users. The focus is not only translation accuracy, but also the organization of information: what the service does, what it does not do, what the client must send and how revisions are handled."),
                ("h1", "Spanish adaptation"),
                ("p", "Esta muestra analiza cómo la documentación bilingüe de servicios puede aumentar la comprensión en organizaciones pequeñas que trabajan con usuarios internacionales. El enfoque no se limita a la precisión de la traducción, sino también a la organización de la información: qué hace el servicio, qué no hace, qué debe enviar el cliente y cómo se gestionan las revisiones."),
                ("h1", "Terminology control"),
                ("table", ["Term", "Preferred rendering", "Note"], [
                    ["scope", "alcance", "Use for service limits and deliverables."],
                    ["revision round", "ronda de revisión", "Avoid vague wording such as correction."],
                    ["source language", "idioma origen", "Keep paired with target language."],
                    ["target language", "idioma destino", "Use consistently in forms and quotes."],
                    ["confidential handling", "manejo confidencial", "Use before receiving sensitive files."],
                ]),
                ("h1", "Editorial notes"),
                ("bullets", [
                    "The abstract is shortened before translation to avoid inflated academic phrasing.",
                    "Key service terms are standardized before drafting the final version.",
                    "The final text keeps a professional register without promising academic acceptance.",
                    "If the client requests a journal style, that requirement is quoted separately.",
                ]),
            ],
        },
        {
            "slug": "scriptorium-classical-annotation-pack",
            "title": "Classical Annotation and Transliteration Pack",
            "subtitle": "Original educational sample for historical-language presentation and study notes.",
            "cover_title": "Classical Annotation",
            "cover_subtitle": "Transliteration, sample translation and cultural notes",
            "code": "CA",
            "metadata": [
                ("Document type", "Original ancient-language educational sample"),
                ("Use", "Annotated samples, study guides, script notes and cultural documentation"),
                ("Commercial status", "Quote-only; educational annotation, not sworn translation"),
                ("Languages", "Latin, Ancient Greek, Hebrew, Classical Arabic, Sanskrit, Old Norse; modern explanation in Spanish and English."),
            ],
            "sections": [
                ("h1", "Purpose"),
                ("p", "Ancient-language work should be presented with humility and structure. This sample separates original form, transliteration, literal meaning and explanatory note so the reader can see what is known, what is interpreted and what remains contextual."),
                ("h1", "Annotated examples"),
                ("table", ["Language", "Original", "Transliteration", "Study note"], [
                    ["Latin", "Vita brevis, ars longa", "vita brevis, ars longa", "A concise aphoristic structure often used to discuss art, skill and time."],
                    ["Ancient Greek", "γνῶθι σεαυτόν", "gnothi seauton", "Imperative phrase traditionally rendered as 'know yourself'."],
                    ["Hebrew", "שָׁלוֹם", "shalom", "Can point to peace, wholeness or wellbeing depending on context."],
                    ["Classical Arabic", "علم", "ilm", "Knowledge; root-based interpretation depends on vocalization and context."],
                    ["Sanskrit", "धर्म", "dharma", "Duty, order, law or teaching depending on tradition and passage."],
                    ["Old Norse", "ᚠᚢᚦᚨᚱᚲ", "futhark", "Runic sequence used as an entry point for script study."],
                ]),
                ("h1", "Delivery model"),
                ("bullets", [
                    "Short excerpt with original script where possible.",
                    "Transliteration using a readable system agreed before delivery.",
                    "Literal meaning separated from interpretive explanation.",
                    "Notes about uncertainty, tradition and educational limits.",
                ]),
                ("h1", "Professional limit"),
                ("p", "This work is presented as study, annotation and educational documentation. It is not legal, religious, medical or certified translation advice."),
            ],
        },
    ]

    for sample in samples:
        save_docx(f"{sample['slug']}.docx", sample["title"], sample["subtitle"], sample["metadata"], sample["sections"])
        save_pdf(f"{sample['slug']}.pdf", sample["title"], sample["subtitle"], sample["metadata"], sample["sections"])
        create_cover(f"{sample['slug']}-preview.png", sample["cover_title"], sample["cover_subtitle"], sample["code"])


if __name__ == "__main__":
    build_all()
