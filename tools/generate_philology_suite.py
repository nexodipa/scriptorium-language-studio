from __future__ import annotations

import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
PORTFOLIO = ROOT / "portfolio"
QA = ROOT / "qa" / "philology-suite"

INK = "#172032"
INK_SOFT = "#314158"
MUTED = "#667386"
BLUE = "#476a78"
BLUE_DEEP = "#243f50"
SAGE = "#8fa49a"
GOLD = "#c9a857"
IVORY = "#fbf7ee"
STONE = "#edf3f4"
LINE = "#d7e0e4"


def font(name: str, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = {
        "serif": [
            r"C:\Windows\Fonts\georgia.ttf",
            r"C:\Windows\Fonts\georgiab.ttf",
            r"C:\Windows\Fonts\times.ttf",
        ],
        "sans": [
            r"C:\Windows\Fonts\segoeui.ttf",
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
        ],
        "sans-bold": [
            r"C:\Windows\Fonts\segoeuib.ttf",
            r"C:\Windows\Fonts\arialbd.ttf",
            r"C:\Windows\Fonts\calibrib.ttf",
        ],
    }
    for item in candidates.get(name, candidates["sans"]):
        if Path(item).exists():
            return ImageFont.truetype(item, size=size)
    return ImageFont.load_default()


def hex_to_rgba(hex_color: str, alpha: int = 255) -> tuple[int, int, int, int]:
    value = hex_color.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4)) + (alpha,)


def cover_crop(image: Image.Image, size: tuple[int, int], anchor: str = "center") -> Image.Image:
    target_w, target_h = size
    source_w, source_h = image.size
    scale = max(target_w / source_w, target_h / source_h)
    resized = image.resize((int(source_w * scale), int(source_h * scale)), Image.Resampling.LANCZOS)
    left = (resized.width - target_w) // 2
    if anchor == "right":
        left = resized.width - target_w
    if anchor == "left":
        left = 0
    top = (resized.height - target_h) // 2
    return resized.crop((left, top, left + target_w, top + target_h))


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    max_width: int,
    font_obj: ImageFont.ImageFont,
    fill: str,
    line_gap: int = 10,
    spacing_before: int = 0,
) -> int:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=font_obj)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)

    x, y = xy
    y += spacing_before
    for line in lines:
        draw.text((x, y), line, font=font_obj, fill=fill)
        bbox = draw.textbbox((x, y), line, font=font_obj)
        y += bbox[3] - bbox[1] + line_gap
    return y


def draw_brand_bar(draw: ImageDraw.ImageDraw, width: int, y: int, label: str) -> None:
    draw.line((72, y, width - 72, y), fill=hex_to_rgba(GOLD, 180), width=2)
    draw.text((72, y + 18), "SCRIPTORIUM", font=font("serif", 30), fill=BLUE_DEEP)
    draw.text((72, y + 56), label.upper(), font=font("sans-bold", 15), fill=MUTED)


def create_hero_derivatives() -> list[Path]:
    base = Image.open(ASSETS / "scriptorium-philology-hero.png").convert("RGB")
    logo = Image.open(ASSETS / "logo-icon.png").convert("RGBA")
    outputs: list[Path] = []

    hero = cover_crop(base, (1800, 980), anchor="center").convert("RGBA")
    veil = Image.new("RGBA", hero.size, (255, 255, 255, 0))
    veil_draw = ImageDraw.Draw(veil)
    for x in range(hero.width):
        alpha = int(210 * (1 - x / hero.width))
        veil_draw.line((x, 0, x, hero.height), fill=(251, 247, 238, max(alpha, 0)))
    hero = Image.alpha_composite(hero, veil)
    out = ASSETS / "scriptorium-philology-hero-framed.png"
    hero.convert("RGB").save(out, quality=94)
    outputs.append(out)

    layers = Image.new("RGB", (1600, 1040), IVORY)
    draw = ImageDraw.Draw(layers)
    for i in range(0, 1600, 40):
        draw.line((i, 0, i, 1040), fill="#eef0eb")
    for i in range(0, 1040, 40):
        draw.line((0, i, 1600, i), fill="#eef0eb")
    draw.rectangle((70, 70, 1530, 970), outline=LINE, width=2)
    draw.text((110, 110), "Mapa de trabajo filológico", font=font("serif", 64), fill=INK)
    draw_wrapped(
        draw,
        "Una lectura profesional no avanza en línea recta: separa forma, lengua, contexto, traducción y entrega para que el texto llegue claro al lector final.",
        (112, 204),
        1050,
        font("sans", 31),
        INK_SOFT,
        line_gap=12,
    )
    steps = [
        ("01", "Texto", "archivo, soporte, estado, legibilidad"),
        ("02", "Forma", "grafía, estructura, citas, formato"),
        ("03", "Lengua", "gramática, registro, terminología"),
        ("04", "Contexto", "época, género, lector, uso"),
        ("05", "Entrega", "traducción, nota, glosario, publicación"),
    ]
    x = 115
    y = 402
    for number, title, desc in steps:
        draw.rounded_rectangle((x, y, x + 255, y + 320), radius=16, fill="#ffffff", outline=LINE, width=2)
        draw.text((x + 26, y + 28), number, font=font("serif", 46), fill=GOLD)
        draw.text((x + 26, y + 102), title, font=font("serif", 40), fill=BLUE_DEEP)
        draw_wrapped(draw, desc, (x + 28, y + 166), 190, font("sans", 24), MUTED, line_gap=8)
        if number != "05":
            draw.line((x + 266, y + 160, x + 306, y + 160), fill=GOLD, width=3)
        x += 292
    out = ASSETS / "scriptorium-philology-layers.png"
    layers.save(out, quality=94)
    outputs.append(out)

    social_specs = [
        ("social-facebook-cover.png", (1640, 624), "Filología aplicada", "Traducción, localización y documentación lingüística por cotización.", "right"),
        ("social-instagram-post.png", (1080, 1080), "Precision across languages", "Textos modernos, lenguas antiguas y entregables profesionales.", "center"),
        ("social-instagram-story.png", (1080, 1920), "Scriptorium", "Servicios de idioma con método filológico, criterio editorial y entrega clara.", "center"),
    ]
    for filename, size, title, subtitle, anchor in social_specs:
        canvas_img = cover_crop(base, size, anchor=anchor).convert("RGBA")
        overlay = Image.new("RGBA", size, (0, 0, 0, 0))
        d = ImageDraw.Draw(overlay)
        for yy in range(size[1]):
            alpha = int(205 * (1 - yy / size[1])) if size[1] > size[0] else 168
            d.line((0, yy, size[0], yy), fill=(251, 247, 238, min(max(alpha, 70), 220)))
        canvas_img = Image.alpha_composite(canvas_img, overlay)
        d = ImageDraw.Draw(canvas_img)
        pad = 84 if size[0] > 1200 else 64
        logo_w = 118 if size[0] > 1200 else 92
        logo_resized = logo.resize((logo_w, logo_w), Image.Resampling.LANCZOS)
        canvas_img.alpha_composite(logo_resized, (pad, pad))
        d.text((pad, pad + logo_w + 42), "SCRIPTORIUM", font=font("serif", 64 if size[0] > 1200 else 52), fill=INK)
        d.text((pad, pad + logo_w + 124), title.upper(), font=font("sans-bold", 25 if size[0] > 1200 else 22), fill=BLUE_DEEP)
        draw_wrapped(d, subtitle, (pad, pad + logo_w + 182), min(720, size[0] - pad * 2), font("sans", 30 if size[0] > 1200 else 28), INK_SOFT, line_gap=12)
        d.line((pad, size[1] - pad - 46, size[0] - pad, size[1] - pad - 46), fill=hex_to_rgba(GOLD, 190), width=3)
        d.text((pad, size[1] - pad - 30), "josuepug@gmail.com  |  WhatsApp +593 98 741 1592", font=font("sans-bold", 19 if size[0] > 1200 else 17), fill=BLUE_DEEP)
        out = ASSETS / filename
        canvas_img.convert("RGB").save(out, quality=94)
        outputs.append(out)

    return outputs


def create_preview_card(filename: str, title: str, label: str, description: str) -> Path:
    base = Image.open(ASSETS / "scriptorium-philology-hero.png").convert("RGB")
    card = cover_crop(base, (1400, 900), anchor="right").convert("RGBA")
    overlay = Image.new("RGBA", card.size, (251, 247, 238, 0))
    d = ImageDraw.Draw(overlay)
    d.rectangle((0, 0, 870, 900), fill=(251, 247, 238, 236))
    card = Image.alpha_composite(card, overlay)
    d = ImageDraw.Draw(card)
    d.rectangle((64, 64, 1336, 836), outline=hex_to_rgba(GOLD, 150), width=2)
    d.text((112, 114), "SCRIPTORIUM", font=font("serif", 42), fill=BLUE_DEEP)
    d.text((112, 172), label.upper(), font=font("sans-bold", 18), fill=MUTED)
    draw_wrapped(d, title, (112, 270), 640, font("serif", 58), INK, line_gap=10)
    draw_wrapped(d, description, (112, 522), 650, font("sans", 29), INK_SOFT, line_gap=12)
    d.line((112, 742, 710, 742), fill=hex_to_rgba(GOLD, 190), width=3)
    out = ASSETS / filename
    card.convert("RGB").save(out, quality=94)
    return out


def setup_doc(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, side, Inches(1))

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(23, 32, 50)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def add_title(document: Document, title: str, subtitle: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.font.name = "Georgia"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("172032")
    run.bold = True
    p.paragraph_format.space_after = Pt(3)

    p = document.add_paragraph()
    run = p.add_run(subtitle)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("667386")
    p.paragraph_format.space_after = Pt(14)


def add_callout(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string("243F50")
    para.paragraph_format.space_after = Pt(0)
    document.add_paragraph()


def build_docx(title: str, subtitle: str, sections: list[tuple[str, list[str]]], out_path: Path, include_template_table: bool = False) -> None:
    document = Document()
    setup_doc(document)
    add_title(document, title, subtitle)
    document.add_picture(str(ASSETS / "scriptorium-philology-layers.png"), width=Inches(6.5))
    caption = document.add_paragraph("Visual de referencia: capas de trabajo filológico de Scriptorium.")
    caption.runs[0].font.size = Pt(8.5)
    caption.runs[0].font.color.rgb = RGBColor.from_string("667386")

    add_callout(
        document,
        "Uso sugerido: enviar este archivo como referencia comercial, guía interna o material previo a una cotización. No sustituye certificación juramentada ni asesoría legal.",
    )

    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for item in paragraphs:
            if item.startswith("LIST:"):
                for bullet in item[5:].split("|"):
                    document.add_paragraph(bullet.strip(), style="List Bullet")
            else:
                document.add_paragraph(item)

    if include_template_table:
        document.add_heading("Brief editable para cotización", level=1)
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width = Inches(1.85)
        table.columns[1].width = Inches(4.65)
        table.rows[0].cells[0].text = "Campo"
        table.rows[0].cells[1].text = "Respuesta"
        for label in [
            "Nombre / organizacion",
            "Canal de contacto",
            "Idioma origen",
            "Idioma destino",
            "Tipo de texto",
            "Numero aproximado de palabras",
            "Formato de archivo",
            "Fecha deseada",
            "Nivel de revision",
            "Confidencialidad",
            "Uso final del material",
            "Notas terminologicas",
        ]:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = ""

    document.core_properties.author = "Scriptorium"
    document.core_properties.title = title
    document.save(out_path)


def wrap_pdf_text(text: str, max_width: float, font_name: str, size: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if stringWidth(trial, font_name, size) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def build_pdf(title: str, subtitle: str, sections: list[tuple[str, list[str]]], out_path: Path, include_template_table: bool = False) -> None:
    c = canvas.Canvas(str(out_path), pagesize=letter)
    width, height = letter
    margin = inch

    def header() -> float:
        c.setFillColor(colors.HexColor(BLUE_DEEP))
        c.setFont("Times-Bold", 24)
        c.drawString(margin, height - margin, "Scriptorium")
        c.setFillColor(colors.HexColor(GOLD))
        c.setFont("Helvetica-Bold", 8)
        c.drawString(margin, height - margin - 18, subtitle.upper()[:90])
        c.setStrokeColor(colors.HexColor(GOLD))
        c.line(margin, height - margin - 32, width - margin, height - margin - 32)
        return height - margin - 62

    y = header()
    c.setFillColor(colors.HexColor(INK))
    c.setFont("Times-Bold", 21)
    for line in wrap_pdf_text(title, width - 2 * margin, "Times-Bold", 21):
        c.drawString(margin, y, line)
        y -= 25
    y -= 12
    c.setFillColor(colors.HexColor(INK_SOFT))
    c.setFont("Helvetica", 10.5)
    lead = "Documento de referencia para presentar el enfoque de filología aplicada, servicios por cotización y entregables profesionales de Scriptorium."
    for line in wrap_pdf_text(lead, width - 2 * margin, "Helvetica", 10.5):
        c.drawString(margin, y, line)
        y -= 15
    y -= 10

    for heading, paragraphs in sections:
        if y < 1.8 * inch:
            c.showPage()
            y = header()
        c.setFillColor(colors.HexColor(BLUE_DEEP))
        c.setFont("Helvetica-Bold", 13)
        c.drawString(margin, y, heading)
        y -= 19
        c.setFillColor(colors.HexColor(MUTED))
        c.setFont("Helvetica", 9.8)
        for item in paragraphs:
            parts = item[5:].split("|") if item.startswith("LIST:") else [item]
            for part in parts:
                prefix = "- " if item.startswith("LIST:") else ""
                for line in wrap_pdf_text(prefix + part.strip(), width - 2 * margin, "Helvetica", 9.8):
                    if y < 0.9 * inch:
                        c.showPage()
                        y = header()
                        c.setFillColor(colors.HexColor(MUTED))
                        c.setFont("Helvetica", 9.8)
                    c.drawString(margin, y, line)
                    y -= 14
                y -= 3
        y -= 10

    if include_template_table:
        if y < 3.0 * inch:
            c.showPage()
            y = header()
        c.setFillColor(colors.HexColor(BLUE_DEEP))
        c.setFont("Helvetica-Bold", 13)
        c.drawString(margin, y, "Campos del brief de servicio")
        y -= 24
        c.setFont("Helvetica", 9)
        for label in ["Cliente", "Contacto", "Idioma origen", "Idioma destino", "Tipo de texto", "Palabras", "Formato", "Fecha", "Revisión", "Confidencialidad"]:
            c.setStrokeColor(colors.HexColor(LINE))
            c.rect(margin, y - 7, width - 2 * margin, 24, stroke=1, fill=0)
            c.drawString(margin + 8, y, label)
            y -= 24

    c.setFillColor(colors.HexColor(MUTED))
    c.setFont("Helvetica", 8)
    c.drawRightString(width - margin, 0.45 * inch, "Scriptorium - documento de referencia")
    c.save()


def validate_pdf(path: Path) -> int:
    reader = PdfReader(str(path))
    return len(reader.pages)


def main() -> None:
    PORTFOLIO.mkdir(exist_ok=True)
    ASSETS.mkdir(exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)

    asset_outputs = create_hero_derivatives()
    asset_outputs.extend(
        [
            create_preview_card(
                "scriptorium-philology-profile-preview.png",
                "Perfil de estudio filológico",
                "Documento comercial",
                "Presenta la identidad, enfoque, límites y entregables de Scriptorium.",
            ),
            create_preview_card(
                "scriptorium-method-guide-preview.png",
                "Guía de método aplicado",
                "Guía profesional",
                "Ordena lectura, traducción, localización, revisión y publicación.",
            ),
            create_preview_card(
                "scriptorium-social-kit-preview.png",
                "Kit de lanzamiento social",
                "Facebook e Instagram",
                "Lineamientos de bio, posts iniciales, tono y piezas visuales.",
            ),
            create_preview_card(
                "scriptorium-service-brief-template-preview.png",
                "Plantilla de brief de servicio",
                "DOCX editable",
                "Campos listos para cotizar proyectos de idioma con más precisión.",
            ),
        ]
    )

    docs = [
        (
            "scriptorium-philology-studio-profile",
            "Perfil de estudio filológico aplicado",
            "Estudio de idiomas, textos antiguos, documentación y publicaciones",
            [
                (
                    "Qué posiciona a Scriptorium",
                    [
                        "Scriptorium se presenta como un estudio de idiomas con método filológico aplicado: no vende solo conversión palabra por palabra, sino lectura, contexto, forma documental y entrega profesional.",
                        "El enfoque combina traducción, localización, apoyo editorial, documentación académica, materiales educativos y notas de lenguas históricas bajo cotización previa.",
                    ],
                ),
                (
                    "Promesa operativa",
                    [
                        "Cada solicitud se evalúa por idioma, extensión, formato, uso final, urgencia, sensibilidad y nivel de revisión. Esto protege al cliente y evita promesas genéricas.",
                        "LIST: Propuestas claras por alcance|Entregables editables cuando corresponda|Uso prudente de muestras y fragmentos|Separación entre trabajo educativo y certificaciones oficiales",
                    ],
                ),
                (
                    "Límites profesionales",
                    [
                        "Scriptorium no se presenta como traducción juramentada, despacho legal ni certificador académico. Los documentos sensibles requieren acuerdo previo de confidencialidad, alcance y canal de entrega.",
                    ],
                ),
            ],
        ),
        (
            "scriptorium-applied-philology-method-guide",
            "Guía de método filológico aplicado",
            "Lectura, traducción, localización y documentación con criterio editorial",
            [
                (
                    "Las cinco capas del trabajo",
                    [
                        "Texto: estado del archivo, soporte, legibilidad y estructura. Forma: grafía, citas, tablas, interfaz o maquetación. Lengua: gramática, registro y terminología. Contexto: género, época, lector y uso. Entrega: formato final, nota, glosario o publicación.",
                        "El método ayuda a que un proyecto no se pierda entre idiomas, formatos y expectativas distintas.",
                    ],
                ),
                (
                    "Aplicación por tipo de servicio",
                    [
                        "LIST: Traducción documental: fidelidad, claridad y revisión|Localización: usuario, tono, interfaz y conversión|Académico: término, cita, registro y disciplina|Lenguas antiguas: transliteración, nota gramatical y contexto educativo|Publicación: estructura, versión bilingüe y presentación final",
                    ],
                ),
                (
                    "Revisión antes de aceptar",
                    [
                        "Cuando un archivo contiene material confidencial, médico, legal, académico sensible o imágenes de baja calidad, se solicita primero un fragmento o descripción para evaluar viabilidad sin exponer todo el documento.",
                    ],
                ),
            ],
        ),
        (
            "scriptorium-social-launch-kit",
            "Kit de lanzamiento para Facebook e Instagram",
            "Perfil, bio, tono, primeras publicaciones y activos visuales",
            [
                (
                    "Bio sugerida",
                    [
                        "Scriptorium | Filología aplicada, traducción, localización y documentación lingüística. Textos modernos, lenguas antiguas y entregables por cotización. Contacto: josuepug@gmail.com.",
                    ],
                ),
                (
                    "Primeras publicaciones",
                    [
                        "LIST: Presentación de marca: qué es Scriptorium y qué problema resuelve|Método: las cinco capas del trabajo filológico|Servicios: traducción, localización, academia, textos antiguos y publicaciones|Confianza: cómo solicitar cotización sin exponer archivos sensibles|Portafolio: mostrar documentos de referencia y límites profesionales",
                    ],
                ),
                (
                    "Tono de comunicacion",
                    [
                        "El tono debe sonar culto pero accesible: preciso, sereno, concreto y útil. Evitar frases grandilocuentes, promesas absolutas o exceso de misterio. La marca debe sentirse humana, editorial y confiable.",
                    ],
                ),
            ],
        ),
        (
            "scriptorium-service-brief-template",
            "Plantilla de brief de servicio lingüístico",
            "Formato editable para solicitar cotización con información completa",
            [
                (
                    "Cómo usar esta plantilla",
                    [
                        "Completa los campos antes de enviar una solicitud por correo, WhatsApp o Fiverr. Si el material es sensible, describe el alcance primero y espera confirmación de condiciones antes de adjuntar el archivo completo.",
                    ],
                ),
                (
                    "Información mínima",
                    [
                        "LIST: Idioma origen y destino|Tipo de texto o plataforma|Número aproximado de palabras o pantallas|Formato editable disponible|Fecha deseada y nivel de urgencia|Uso final del material|Condiciones de confidencialidad",
                    ],
                ),
            ],
            True,
        ),
    ]

    final_docs: list[Path] = []
    for item in docs:
        slug, title, subtitle, sections, *template = item
        include_template = bool(template and template[0])
        docx_path = PORTFOLIO / f"{slug}.docx"
        pdf_path = PORTFOLIO / f"{slug}.pdf"
        build_docx(title, subtitle, sections, docx_path, include_template_table=include_template)
        build_pdf(title, subtitle, sections, pdf_path, include_template_table=include_template)
        validate_pdf(pdf_path)
        final_docs.extend([docx_path, pdf_path])

    manifest = QA / "generated-files.txt"
    manifest.write_text(
        "\n".join(str(path.relative_to(ROOT)) for path in [*asset_outputs, *final_docs]),
        encoding="utf-8",
    )
    print(f"Generated {len(asset_outputs)} assets and {len(final_docs)} documents.")


if __name__ == "__main__":
    main()
